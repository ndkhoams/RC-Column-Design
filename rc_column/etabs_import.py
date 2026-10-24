"""Đọc nội lực cột từ file Excel do ETABS xuất (bảng Element Forces - Columns).

Cột cần thiết: Story, Column, Output Case, Case Type, Station, P, M2, M3.
Quy ước: P âm = nén (ETABS) -> đổi sang N dương = nén.
Ánh xạ mô men: mặc định M2 -> Mx, M3 -> My (đổi được trên GUI).

Hỗ trợ thêm SAP2000 (bảng Element Forces - Frames): header dùng Frame,
OutputCase, CaseType, StepType (không khoảng trắng) và KHÔNG có cột Story —
khi đó Label lấy từ Frame, Story mặc định "Story". Quy ước dấu/mô men như ETABS.
Clipboard SAP2000 cũng có thể chỉ gồm các dòng số liệu (KHÔNG có dòng header,
khi user chỉ chọn vùng dữ liệu trong bảng SAP) — trường hợp đó parse theo
VỊ TRÍ cột với kiểm chứng chặt (xem _parse_sap_clipboard_noheader).
"""

from dataclasses import dataclass
from typing import List, Optional

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


@dataclass
class EtabsForce:
    story: str
    column: str
    case: str
    case_type: str
    station: float
    N: float    # kN, nén dương
    Mx: float   # kN·m
    My: float   # kN·m
    Q2: float = 0.0  # kN, lực cắt V2
    Q3: float = 0.0  # kN, lực cắt V3


def read_etabs_excel(path: str,
                     only_combinations: bool = True) -> List[EtabsForce]:
    """Đọc file Excel ETABS.

    Quy ước trục 2/3 thống nhất: M2 (Excel) -> M2, M3 (Excel) -> M3.
    """
    if not HAS_OPENPYXL:
        raise ImportError("Cần cài openpyxl: pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True)
    # tìm sheet chứa "Element Forces"
    ws = None
    for name in wb.sheetnames:
        if "element forces" in name.lower():
            ws = wb[name]
            break
    if ws is None:
        ws = wb.active

    # tìm dòng header (ETABS: "Story" + "Output Case";
    # SAP2000: "Frame" + "OutputCase" — không có Story)
    header_row = None
    col_idx = {}
    for r in range(1, min(10, ws.max_row + 1)):
        vals = [str(ws.cell(r, c).value or "").strip().lower()
                for c in range(1, ws.max_column + 1)]
        if ("story" in vals or "frame" in vals) and \
                ("output case" in vals or "outputcase" in vals):
            header_row = r
            for c, v in enumerate(vals, start=1):
                col_idx[v] = c
            break
    if header_row is None:
        raise ValueError("Không tìm thấy bảng Element Forces trong file.")

    def val(row, *names):
        for n in names:
            c = col_idx.get(n)
            if c:
                v = ws.cell(row, c).value
                if v is not None and str(v).strip() != "":
                    return v
        return None

    def num(row, *names):
        v = val(row, *names)
        try:
            return float(v)
        except (TypeError, ValueError):
            return 0.0

    # Tự động nhận diện đơn vị từ dòng ngay sau header (vd "kN", "kN-m")
    def _unit_factor(unit_str: str, kind: str) -> float:
        u = (str(unit_str or "")).strip().lower().replace(" ", "")
        if kind == "force":
            if u in ("kn",): return 1.0
            if u in ("n",): return 0.001
            if u in ("kgf", "kg", "kgs"): return 0.00980665
            if u in ("tonf", "ton", "t", "tf"): return 9.80665
            if u in ("kip", "kips"): return 4.4482216
            if u in ("lbf", "lb", "lbs"): return 0.0044482216
            return 1.0
        else:
            if u in ("kn-m", "knm"): return 1.0
            if u in ("n-m", "nm"): return 0.001
            if u in ("n-mm", "nmm"): return 1e-6
            if u in ("kn-mm", "knmm"): return 0.001
            if u in ("kgf-m", "kg-m", "kgm"): return 0.00980665
            if u in ("kgf-cm", "kg-cm", "kgcm"): return 0.00980665 * 0.01
            if u in ("tonf-m", "ton-m", "tf-m"): return 9.80665
            if u in ("kip-ft", "kipft", "k-ft"): return 1.3558179
            if u in ("kip-in", "kipin", "k-in"): return 1.3558179 / 12.0
            if u in ("lbf-ft", "lb-ft"): return 0.0013558179
            if u in ("lbf-in", "lb-in"): return 0.0013558179 / 12.0
            return 1.0

    f_force, f_moment = 1.0, 1.0
    ur = header_row + 1
    if ur <= ws.max_row:
        c_p = col_idx.get("p")
        c_m2 = col_idx.get("m2")
        if c_p:
            f_force = _unit_factor(ws.cell(ur, c_p).value, "force")
        if c_m2:
            f_moment = _unit_factor(ws.cell(ur, c_m2).value, "moment")

    forces: List[EtabsForce] = []
    data_start = header_row + 2  # bỏ qua dòng đơn vị
    for r in range(data_start, ws.max_row + 1):
        case = val(r, "output case", "outputcase")
        if case is None:
            continue
        case_type = str(val(r, "case type", "casetype") or "")
        if only_combinations and "combination" not in case_type.lower():
            continue
        # Gắn Step Type vào tên tổ hợp: TH14max, TH15min
        step = str(val(r, "step type", "steptype") or "").strip().lower()
        if step == "max":
            case = f"{case}max"
        elif step == "min":
            case = f"{case}min"
        p = num(r, "p") * f_force
        m2 = num(r, "m2") * f_moment
        m3 = num(r, "m3") * f_moment
        v2 = num(r, "v2", "shear 2-2", "v2-2") * f_force
        v3 = num(r, "v3", "shear 3-3", "v3-3") * f_force
        mx, my = m2, m3  # M2->M2, M3->M3 (không ánh xạ)
        forces.append(EtabsForce(
            # SAP2000 không có cột Story -> mặc định "Story" (user đã duyệt)
            story=str(val(r, "story") or
                      ("Story" if "story" not in col_idx else "")),
            column=str(val(r, "column", "frame") or ""),
            case=str(case),
            case_type=case_type,
            station=num(r, "station"),
            N=-p,  # ETABS: nén âm -> đổi sang nén dương
            Mx=mx,
            My=my,
            Q2=v2,
            Q3=v3,
        ))
    return forces


def _parse_sap_clipboard(lines, hi, idx, only_combinations,
                         unit_factor) -> List[EtabsForce]:
    """Parse các dòng dữ liệu SAP2000 (Element Forces - Frames) theo TÊN cột.

    idx: dict tên trường -> vị trí cột trong header. SAP không có Story ->
    story mặc định "Story", Label lấy từ Frame. Quy ước dấu như ETABS
    (N = -P; M2 -> Mx, M3 -> My; hậu tố Step Type gắn vào tên tổ hợp).
    Dòng đơn vị (nếu có) tự bị bỏ qua vì ô P của nó không phải số.
    """
    ip, im2 = idx["p"], idx["m2"]
    f_force, f_moment = 1.0, 1.0
    if hi + 1 < len(lines):
        up = lines[hi + 1].split("\t")
        if ip is not None and ip < len(up):
            f_force = unit_factor(up[ip], "force")
        if im2 is not None and im2 < len(up):
            f_moment = unit_factor(up[im2], "moment")

    def _num(parts, i):
        if i is None or i >= len(parts):
            return 0.0
        try:
            return float(parts[i].strip())
        except ValueError:
            return 0.0

    def _txt(parts, i):
        return parts[i].strip() if i is not None and i < len(parts) else ""

    forces: List[EtabsForce] = []
    for ln in lines[hi + 1:]:
        parts = ln.split("\t")
        if ip is None or len(parts) <= ip:
            continue
        try:
            p_raw = float(parts[ip].strip())
        except (ValueError, AttributeError):
            continue  # dòng đơn vị / dòng rác: ô P không phải số
        case = _txt(parts, idx["case"])
        if not case:
            continue
        case_type = _txt(parts, idx["type"])
        if only_combinations and "combination" not in case_type.lower():
            continue
        # Gắn Step Type vào tên tổ hợp như ETABS: ENV SLS + Min -> ENV SLSmin
        step = _txt(parts, idx["step"]).lower()
        if step == "max":
            case = f"{case}max"
        elif step == "min":
            case = f"{case}min"
        forces.append(EtabsForce(
            story="Story",
            column=_txt(parts, idx["label"]),
            case=case,
            case_type=case_type,
            station=_num(parts, idx["station"]),
            N=-p_raw * f_force,
            Mx=_num(parts, idx["m2"]) * f_moment,
            My=_num(parts, idx["m3"]) * f_moment,
            Q2=_num(parts, idx["v2"]) * f_force,
            Q3=_num(parts, idx["v3"]) * f_force,
        ))
    if not forces:
        raise ValueError(
            "Không đọc được dòng nội lực nào từ clipboard SAP2000.")
    return forces


# Case Type SAP2000 có thể xuất trong bảng Element Forces - Frames — dùng
# để NHẬN DIỆN clipboard SAP không header (kiểm chứng chặt, tránh parse nhầm
# dữ liệu rác / layout khác thành nội lực).
_SAP_CASE_TYPES = {
    "combination", "linstatic", "linrespspec", "linreshist", "linmodal",
    "linbuckling", "nonstatic", "envelope",
}
_SAP_STEP_TYPES = {"", "max", "min", "absolute"}


def _is_number(s: str) -> bool:
    try:
        float(s)
        return True
    except (TypeError, ValueError):
        return False


def _match_sap_noheader_row(parts):
    """Kiểm chứng 1 dòng clipboard SAP2000 KHÔNG header theo vị trí cột.

    Layout: Frame | Station | OutputCase | CaseType | [StepType] | P | V2 |
    V3 | T | M2 | M3 | [FrameElem | ElemStation].
    Trả về (p_idx, step) nếu dòng KHỚP CHẶT, ngược lại None:
    - Station phải là số; OutputCase phải là chữ (không phải số);
      CaseType thuộc tập Case Type SAP2000 đã biết.
    - Có cột StepType khi trường ngay sau CaseType KHÔNG phải số mà thuộc
      {rỗng, Max, Min, Absolute} và trường kế tiếp là số (chính là P).
    - 6 trường P..M3 liên tiếp phải đều là số; phần đuôi sau M3 chỉ được
      có 0 hoặc 2 trường, và ElemStation (trường cuối) phải là số.
    """
    parts = [p.strip() for p in parts]
    if len(parts) < 10:
        return None
    frame, station, case, case_type = parts[0], parts[1], parts[2], parts[3]
    if not frame or not case or _is_number(case):
        return None
    if not _is_number(station):
        return None
    if case_type.lower().replace(" ", "") not in _SAP_CASE_TYPES:
        return None
    if _is_number(parts[4]):
        p_idx, step = 4, ""
    elif (parts[4].lower() in _SAP_STEP_TYPES and len(parts) >= 11
            and _is_number(parts[5])):
        p_idx, step = 5, parts[4].lower()
    else:
        return None
    if len(parts) < p_idx + 6:
        return None
    if not all(_is_number(x) for x in parts[p_idx:p_idx + 6]):
        return None
    trail = parts[p_idx + 6:]
    if len(trail) not in (0, 2):
        return None
    if len(trail) == 2 and not _is_number(trail[1]):
        return None
    return p_idx, step


def _parse_sap_clipboard_noheader(lines,
                                  only_combinations) -> List[EtabsForce]:
    """Parse clipboard SAP2000 chỉ gồm các dòng số liệu (KHÔNG có header).

    Chỉ kích hoạt khi dòng đầu tiên khớp chặt layout SAP theo vị trí
    (xem _match_sap_noheader_row); mọi dòng sau phải cùng layout — dòng
    nào lệch sẽ báo lỗi kèm SỐ DÒNG thay vì parse sai thầm lặng. Nếu
    dòng đầu không khớp thì giữ nguyên lỗi "Không tìm thấy header".
    Ánh xạ như đường SAP có header: Story = "Story", Label <- Frame,
    N = -P, Mx <- M2, My <- M3, hậu tố Step Type gắn vào tên tổ hợp.
    Không có dòng đơn vị nên giữ nguyên số (đơn vị kN, kN-m như bản
    xuất SAP2000 mặc định).
    """
    first = _match_sap_noheader_row(lines[0].split("\t"))
    if first is None:
        raise ValueError("Không tìm thấy header ETABS (Output Case).")
    p_idx = first[0]
    n_fields = len(lines[0].split("\t"))

    forces: List[EtabsForce] = []
    for i, ln in enumerate(lines):
        parts = ln.split("\t")
        m = _match_sap_noheader_row(parts)
        if m is None or len(parts) != n_fields or m[0] != p_idx:
            raise ValueError(
                f"Dòng {i + 1} trong clipboard không đúng định dạng "
                f"SAP2000 (Element Forces - Frames không header).")
        vals = [float(x) for x in parts[p_idx:p_idx + 6]]  # P,V2,V3,T,M2,M3
        case = parts[2].strip()
        case_type = parts[3].strip()
        if only_combinations and "combination" not in case_type.lower():
            continue
        if m[1] == "max":
            case = f"{case}max"
        elif m[1] == "min":
            case = f"{case}min"
        forces.append(EtabsForce(
            story="Story",
            column=parts[0].strip(),
            case=case,
            case_type=case_type,
            station=float(parts[1]),
            N=-vals[0],
            Mx=vals[4],
            My=vals[5],
            Q2=vals[1],
            Q3=vals[2],
        ))
    if not forces:
        raise ValueError(
            "Không đọc được dòng nội lực nào từ clipboard SAP2000.")
    return forces


def _strip_edge_empty_fields(line: str) -> str:
    """Bỏ các trường RỖNG ở đầu/cuối 1 dòng clipboard (tách theo tab).

    Khi copy từ bảng hiển thị SAP2000/ETABS, mỗi dòng thường dính 1 ô
    rỗng ở ĐẦU dòng (cột số thứ tự dòng của bảng) và có thể dính ô rỗng
    ở CUỐI dòng -> các dòng lệch số trường dù cùng layout, parser đối
    chiếu số trường sẽ chết ở dòng 2. Trường rỗng ở GIỮA dòng (vd Step
    Type rỗng) được GIỮ NGUYÊN.
    """
    parts = line.split("\t")
    while parts and not parts[0].strip():
        parts.pop(0)
    while parts and not parts[-1].strip():
        parts.pop()
    return "\t".join(parts)


def parse_etabs_clipboard(text: str,
                          only_combinations: bool = True) -> List[EtabsForce]:
    """Đọc nội lực từ clipboard copy từ bảng ETABS 'Element Forces - Columns'.

    Định dạng: tab-separated, 2 dòng header, các cột:
    Story | Column | Unique Name | Output Case | Case Type | [Step Type] |
    Station | P | V2 | V3 | T | M2 | M3 | Element | Elem Station | Location
    (tự động nhận diện có/không có cột Step Type để match đúng vị trí)
    Quy ước: P âm = nén (ETABS) -> N dương.

    Nhận thêm clipboard SAP2000 (bảng Element Forces - Frames): header theo
    tên cột Frame | Station | OutputCase | CaseType | StepType | P | V2 | V3 |
    T | M2 | M3 | ... (không có Story) — parse theo TÊN cột, Story = "Story".
    Clipboard SAP2000 KHÔNG có header (chỉ các dòng số liệu) cũng được nhận:
    parse theo VỊ TRÍ cột với kiểm chứng chặt từng dòng
    (xem _parse_sap_clipboard_noheader).
    """
    # Chuẩn hoá từng dòng TRƯỚC mọi định tuyến/kiểm chứng: bỏ ô rỗng ở
    # đầu/cuối dòng (tab thừa khi copy từ bảng — xem _strip_edge_empty_fields)
    lines = [_strip_edge_empty_fields(ln)
             for ln in text.strip().splitlines() if ln.strip()]
    if len(lines) < 3:
        raise ValueError("Clipboard không có đủ dữ liệu.")
    # Tìm dòng header chứa "Output Case" (ETABS) hoặc "OutputCase" (SAP2000)
    hi = None
    for i, ln in enumerate(lines[:5]):
        low = ln.lower()
        if "output case" in low or "outputcase" in low:
            hi = i
            break
    if hi is None:
        # Không có dòng header: thử clipboard SAP2000 chỉ gồm các dòng
        # số liệu (parse theo vị trí với kiểm chứng chặt). Nếu dòng đầu
        # không khớp layout SAP, hàm này tự ném lại đúng lỗi
        # "Không tìm thấy header ETABS (Output Case)." như trước.
        return _parse_sap_clipboard_noheader(lines, only_combinations)
    # Tự động nhận diện có cột "Step Type" hay không để match đúng vị trí
    # Có Step Type: Story|Column|Unique Name|Output Case|Case Type|Step Type|
    #               Station|P|V2|V3|T|M2|M3|...
    # Không Step Type: Story|Column|Unique Name|Output Case|Case Type|
    #               Station|P|V2|V3|T|M2|M3|...
    header_parts = [p.strip().lower() for p in lines[hi].split("\t")]
    # SAP2000: header có "frame" và KHÔNG có "story" -> parse theo tên cột
    # (vị trí cột khác hẳn ETABS nên không dùng được vị trí cố định bên dưới)
    sap_idx = None
    if "frame" in header_parts and "story" not in header_parts:
        def _find(*names):
            for n in names:
                if n in header_parts:
                    return header_parts.index(n)
            return None
        sap_idx = {
            "label": _find("frame"),
            "station": _find("station"),
            "case": _find("output case", "outputcase"),
            "type": _find("case type", "casetype"),
            "step": _find("step type", "steptype"),
            "p": _find("p"),
            "v2": _find("v2"),
            "v3": _find("v3"),
            "m2": _find("m2"),
            "m3": _find("m3"),
        }
        if sap_idx["case"] is None or sap_idx["p"] is None:
            sap_idx = None
    # Kiểm tra Step Type trên cả dòng header và các dòng lân cận
    # (ETABS đôi khi xuất header nhiều dòng)
    has_step = "step type" in header_parts
    if not has_step:
        for _ln in lines[max(0, hi-2):hi+3]:
            if "step type" in _ln.lower():
                has_step = True
                break
    if has_step:
        C_CASE, C_TYPE, C_STATION = 3, 4, 6
        C_P, C_M2, C_M3 = 7, 11, 12
        C_V2, C_V3 = 8, 9
    else:
        C_CASE, C_TYPE, C_STATION = 3, 4, 5
        C_P, C_M2, C_M3 = 6, 10, 11
        C_V2, C_V3 = 7, 8
    C_STORY, C_COL = 0, 1

    # Tự động nhận diện đơn vị từ dòng units (dòng ngay sau header)
    # Quy về kN (lực) và kN-m (mô men) để dùng trong app
    def _unit_factor(unit_str: str, kind: str) -> float:
        """Trả về hệ số nhân để đổi về kN (kind='force') hoặc kN-m (kind='moment')."""
        u = (unit_str or "").strip().lower().replace(" ", "")
        if kind == "force":
            if u in ("kn",): return 1.0
            if u in ("n",): return 0.001
            if u in ("kgf", "kg", "kgs"): return 0.00980665
            if u in ("tonf", "ton", "t", "tf"): return 9.80665
            if u in ("kip", "kips"): return 4.4482216
            if u in ("lbf", "lb", "lbs"): return 0.0044482216
            return 1.0  # mặc định kN
        else:  # moment -> kN-m
            if u in ("kn-m", "knm", "kN-m".lower()): return 1.0
            if u in ("n-m", "nm"): return 0.001
            if u in ("n-mm", "nmm"): return 1e-6
            if u in ("kn-mm", "knmm"): return 0.001
            if u in ("kgf-m", "kg-m", "kgm"): return 0.00980665
            if u in ("kgf-cm", "kg-cm", "kgcm"): return 0.00980665 * 0.01
            if u in ("tonf-m", "ton-m", "tf-m"): return 9.80665
            if u in ("kip-ft", "kipft", "k-ft"): return 1.3558179
            if u in ("kip-in", "kipin", "k-in"): return 1.3558179 / 12.0
            if u in ("lbf-ft", "lb-ft"): return 0.0013558179
            if u in ("lbf-in", "lb-in"): return 0.0013558179 / 12.0
            return 1.0  # mặc định kN-m

    # Nhánh SAP2000: parse theo TÊN cột (vị trí khác hẳn ETABS)
    if sap_idx is not None:
        return _parse_sap_clipboard(lines, hi, sap_idx,
                                    only_combinations, _unit_factor)

    f_force, f_moment = 1.0, 1.0
    if hi + 1 < len(lines):
        unit_parts = lines[hi + 1].split("\t")
        # Lấy đơn vị tại cột P (lực) và cột M2 (mô men)
        u_p = unit_parts[C_P].strip() if C_P < len(unit_parts) else ""
        u_m = unit_parts[C_M2].strip() if C_M2 < len(unit_parts) else ""
        f_force = _unit_factor(u_p, "force")
        f_moment = _unit_factor(u_m, "moment")

    def num(parts, idx):
        try:
            return float(parts[idx].strip()) if idx < len(parts) and parts[idx].strip() else 0.0
        except ValueError:
            return 0.0

    forces: List[EtabsForce] = []
    for ln in lines[hi + 2:]:
        parts = ln.split("\t")
        if len(parts) <= C_P:
            continue
        case = parts[C_CASE].strip() if C_CASE < len(parts) else ""
        if not case:
            continue
        case_type = parts[C_TYPE].strip() if C_TYPE < len(parts) else ""
        if only_combinations and "combination" not in case_type.lower():
            continue
        # Gắn Step Type vào tên tổ hợp: TH14max, TH15min
        if has_step:
            step = parts[5].strip().lower() if len(parts) > 5 else ""
            if step == "max":
                case = f"{case}max"
            elif step == "min":
                case = f"{case}min"
        p = num(parts, C_P) * f_force
        forces.append(EtabsForce(
            story=parts[C_STORY].strip() if C_STORY < len(parts) else "",
            column=parts[C_COL].strip() if C_COL < len(parts) else "",
            case=case,
            case_type=case_type,
            station=num(parts, C_STATION),
            N=-p,
            Mx=num(parts, C_M2) * f_moment,
            My=num(parts, C_M3) * f_moment,
            Q2=num(parts, C_V2) * f_force,
            Q3=num(parts, C_V3) * f_force,
        ))
    if not forces:
        raise ValueError("Không đọc được dòng nội lực nào từ clipboard.")
    return forces


def _max_min_suffix(case: str):
    """Tách hậu tố max/min của ETABS khỏi tên combo.

    ETABS gắn Step Type vào tên tổ hợp ("ULS17max", "ENV SLSmin").
    Trả về (tên gốc, hậu tố "max"/"min") hoặc None nếu tên không có
    hậu tố hợp lệ (tên gốc rỗng cũng tính là không hợp lệ). Hậu tố
    nhận dạng không phân biệt hoa thường; tên gốc giữ nguyên cách viết.
    """
    name = case.strip()
    low = name.lower()
    for suf in ("max", "min"):
        if low.endswith(suf) and len(name) > len(suf):
            return name[:-len(suf)].rstrip(), suf
    return None


def _group_max_min(forces):
    """Nhóm các dòng có hậu tố max/min theo (Story, Label, Station, tên gốc).

    Trả về (groups, order): groups ánh xạ khóa nhóm -> {"max": [idx...],
    "min": [idx...]}; order là thứ tự xuất hiện đầu tiên của các khóa.
    Dòng TX_* và dòng không có hậu tố không vào nhóm.
    """
    groups = {}
    order = []
    for idx, f in enumerate(forces):
        name = f.case.strip()
        if name.startswith("TX_"):
            continue
        split = _max_min_suffix(name)
        if split is None:
            continue
        base, suf = split
        key = (f.story.strip(), f.column.strip(), round(f.station, 6), base)
        if key not in groups:
            groups[key] = {}
            order.append(key)
        groups[key].setdefault(suf, []).append(idx)
    return groups, order


def combo_merge_options(forces):
    """Danh sách tên combo cho popup "Chọn combo động đất…".

    Trả về [(tên gốc, có_cặp, số_cặp)] theo thứ tự xuất hiện đầu tiên
    trong bảng:
    - các biến thể max/min gộp về tên gốc (ULS17max + ULS17min -> "ULS17");
    - tên thường (không hậu tố) liệt kê với có_cặp=False;
    - dòng TX_* loại khỏi danh sách (không bao giờ gộp, không phải combo
      kiểm tra) — popup chỉ liệt kê combo kiểm tra;
    - có_cặp=True khi tồn tại ít nhất một nhóm (Story, Label, Station)
      có đúng 1 dòng max + 1 dòng min của tên gốc đó; số_cặp đếm các
      nhóm như vậy.
    """
    options = []
    pos = {}
    for f in forces:
        name = f.case.strip()
        if name.startswith("TX_"):
            continue
        split = _max_min_suffix(name)
        base = split[0] if split else name
        if base not in pos:
            pos[base] = len(options)
            options.append([base, False, 0])
    groups, _order = _group_max_min(forces)
    for (_story, _col, _st, base), by_suf in groups.items():
        if len(by_suf.get("max", [])) == 1 and len(by_suf.get("min", [])) == 1:
            options[pos[base]][1] = True
            options[pos[base]][2] += 1
    return [(b, hp, n) for b, hp, n in options]


def effective_merge_bases(seismic_selected):
    """Tập tên gốc được gộp khi hiển thị bảng nội lực.

    Tập gộp hiệu dụng = tập động đất: tên gốc được tick trong popup
    "Chọn combo động đất…" thì cặp max/min của nó tự gộp (gộp là hệ
    quả của việc là combo động đất); không còn tập gộp riêng hay nút
    chọn gộp. Bỏ tick thì tên đó không gộp nữa.
    """
    return set(seismic_selected or set())


def merge_combo_max_min(forces, selected_bases=None):
    """Gộp các cặp combo max/min trong bảng nội lực thành 1 dòng.

    Theo chỉ dẫn ETABS về phân tích phổ phản ứng: max/min là biên thống
    kê ± của từng đại lượng riêng lẻ, KHÔNG tương ứng đồng thời giữa các
    thành phần — nên dòng gộp lấy bao theo trị tuyệt đối từng thành phần.

    Quy tắc:
    - Chỉ gộp khi cùng một nhóm (Story, Label, Station) tồn tại cả 2 dòng
      có tên Combo chỉ khác nhau bởi hậu tố max/min (xem _max_min_suffix).
    - selected_bases=None (mặc định): gộp mọi cặp. Nếu cho tập tên gốc:
      chỉ gộp các cặp thuộc tên gốc đó (ngữ nghĩa popup "Chọn combo
      gộp…": tick tên gốc = gộp mọi cặp của tên đó, mọi Story/Station).
    - Dòng gộp: Combo = tên gốc đã bỏ hậu tố; mỗi thành phần (N, V2, V3,
      Mx, My) lấy giá trị có |giá trị| lớn nhất trong 2 dòng, giữ nguyên
      dấu của giá trị đó; hòa trị tuyệt đối thì lấy giá trị của dòng max.
      Story/Label/Station/case_type lấy từ dòng đầu tiên của cặp trong
      bảng gốc; dòng gộp đứng ở vị trí dòng đó.
    - Dòng không có cặp, combo thường và mọi dòng TX_* giữ nguyên (kể cả
      dạng TX_*max/min — bỏ qua cho an toàn).
    - Nhóm có hơn 1 dòng max hoặc hơn 1 dòng min không gộp.
    - Lũy đẳng: dòng đã gộp không còn hậu tố nên gộp lại không đổi gì.
    - Dòng không gộp giữ nguyên OBJECT gốc (để GUI nhận diện dòng gộp
      qua identity khi cần, vd chặn sửa trực tiếp dòng gộp).

    Hàm thuần túy (không chạm GUI), trả về (danh sách mới, số cặp đã gộp).
    """
    groups, _order = _group_max_min(forces)
    merged_row = {}  # vị trí dòng đầu của cặp -> dòng gộp
    consumed = set()  # vị trí các dòng gốc đã gộp (bỏ khỏi bảng mới)
    for key, by_suf in groups.items():
        if selected_bases is not None and key[3] not in selected_bases:
            continue
        maxs = by_suf.get("max", [])
        mins = by_suf.get("min", [])
        if len(maxs) != 1 or len(mins) != 1:
            continue
        f_max = forces[maxs[0]]
        f_min = forces[mins[0]]
        first = forces[min(maxs[0], mins[0])]
        base = _max_min_suffix(first.case)[0]
        vals = {}
        for field in ("N", "Q2", "Q3", "Mx", "My"):
            v_max = getattr(f_max, field)
            v_min = getattr(f_min, field)
            vals[field] = v_max if abs(v_max) >= abs(v_min) else v_min
        merged_row[min(maxs[0], mins[0])] = EtabsForce(
            story=first.story, column=first.column, case=base,
            case_type=first.case_type, station=first.station, **vals)
        consumed.add(maxs[0])
        consumed.add(mins[0])
    out = []
    for idx, f in enumerate(forces):
        if idx in merged_row:
            out.append(merged_row[idx])
        elif idx not in consumed:
            out.append(f)
    return out, len(merged_row)
